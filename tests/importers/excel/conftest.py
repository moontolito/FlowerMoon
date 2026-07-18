from __future__ import annotations

from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import Workbook

from debbie.importers.excel.constants import PART_HEADERS, STOCK_HEADERS, WORK_HEADERS


DEFAULT_WORK = ["W-1", "First Work", 1, 0.2, 0.5, 0.0, 5, 7, 3, 4]
DEFAULT_PART = ["W-1", "P-1", "Panel", 100, 50, 2, "Yes", "DWG-001"]
DEFAULT_STOCK = ["W-1", "S-1", "Sheet", 1000, 500, 2]


@pytest.fixture
def workbook_factory(tmp_path: Path):
    counter = 0

    def create(
        *,
        metadata: list[list[object]] | None = None,
        works: list[list[object]] | None = None,
        parts: list[list[object]] | None = None,
        stocks: list[list[object]] | None = None,
        mutate: Callable[[Workbook], None] | None = None,
        suffix: str = ".xlsx",
    ) -> Path:
        nonlocal counter
        counter += 1
        workbook = Workbook()
        workbook.remove(workbook.active)

        debbie = workbook.create_sheet("Debbie")
        for row in metadata or [
            ["Key", "Value"],
            ["Format Name", "Debbie Nesting Workbook"],
            ["Schema Version", "1.0"],
            ["Units", "mm"],
        ]:
            debbie.append(row)

        work_sheet = workbook.create_sheet("Works")
        work_sheet.append(WORK_HEADERS)
        for row in works if works is not None else [DEFAULT_WORK]:
            work_sheet.append(row)

        part_sheet = workbook.create_sheet("Parts")
        part_sheet.append((*PART_HEADERS, "Drawing Number"))
        for row in parts if parts is not None else [DEFAULT_PART]:
            part_sheet.append(row)

        stock_sheet = workbook.create_sheet("Stocks")
        stock_sheet.append(STOCK_HEADERS)
        for row in stocks if stocks is not None else [DEFAULT_STOCK]:
            stock_sheet.append(row)

        if mutate is not None:
            mutate(workbook)
        path = tmp_path / f"workbook-{counter}{suffix}"
        workbook.save(path)
        workbook.close()
        return path

    return create
