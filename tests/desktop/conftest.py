from __future__ import annotations

import os
from pathlib import Path

import pytest
from openpyxl import Workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication

from debbie.domain import (
    DemandItem,
    DemandItemId,
    Dimensions,
    PartType,
    PartTypeId,
    ProcessProfile,
    StockInstance,
    StockInstanceId,
    StockSpecification,
    StockSpecificationId,
    Trim,
    Work,
    WorkId,
)
from debbie.importers.excel import (
    CanonicalWorkbookImportResult,
    CanonicalWorkbookRecords,
    ImportSummary,
    SourceLocation,
    StockImportRecord,
    WorkbookMetadataRecord,
    WorkImportRecord,
    PartImportRecord,
)
from debbie.importers.excel.identifiers import (
    demand_item_id,
    part_type_id,
    stock_instance_id,
    stock_specification_id,
    work_id,
)

_WORK_KEYS = {}


@pytest.fixture(scope="session")
def qt_app():
    application = QApplication.instance() or QApplication([])
    yield application


def make_work(key: str = "W1", name: str = "Work One", stock_quantity: int = 1) -> Work:
    owner = work_id(key)
    _WORK_KEYS[owner] = key
    type_id = part_type_id(owner, "P-0")
    part_type = PartType(type_id, "Bracket", Dimensions(40, 20))
    demand = DemandItem(demand_item_id(owner, "P-0"), owner, part_type.id, 2)
    specification_id = stock_specification_id(owner, "S-0")
    specification = StockSpecification(
        specification_id, "Standard sheet", Dimensions(100, 60)
    )
    instances = tuple(
        StockInstance(
            stock_instance_id(specification_id, sequence), owner, specification.id, sequence
        )
        for sequence in range(1, stock_quantity + 1)
    )
    return Work.from_inputs(
        id=owner,
        name=name,
        batch_multiplier=2,
        process_profile=ProcessProfile(
            kerf=0.2,
            minimum_part_clearance=1,
            boundary_clearance=2,
            trim=Trim(left=3, right=4, top=5, bottom=6),
        ),
        part_types=(part_type,),
        demand_items=(demand,),
        stock_specifications=(specification,),
        stock_instances=instances,
    )


def make_import_result(*works: Work) -> CanonicalWorkbookImportResult:
    work_records = tuple(
        WorkImportRecord(
            _WORK_KEYS[work.id], work.name, work.batch_multiplier, 0.2, 1, 2, 3, 4, 5, 6,
            SourceLocation("Works", index + 2),
        )
        for index, work in enumerate(works)
    )
    part_records = tuple(
        PartImportRecord(
            _WORK_KEYS[work.id], "P-0", work.part_types[0].name, 40, 20, 2, True,
            f"DWG-{index}", SourceLocation("Parts", index + 2),
        )
        for index, work in enumerate(works)
    )
    stock_records = tuple(
        StockImportRecord(
            _WORK_KEYS[work.id], "S-0", work.stock_specifications[0].name, 100, 60,
            len(work.stock_instances), SourceLocation("Stocks", index + 2),
        )
        for index, work in enumerate(works)
    )
    records = CanonicalWorkbookRecords(
        WorkbookMetadataRecord("Debbie Nesting Workbook", "1.0", "mm"),
        work_records,
        part_records,
        stock_records,
    )
    return CanonicalWorkbookImportResult(
        works=works,
        format_name="Debbie Nesting Workbook",
        schema_version="1.0",
        unit_system="mm",
        source_path=Path("fixture.xlsx"),
        records=records,
        summary=ImportSummary(len(works), len(works), len(works), len(works), len(works)),
    )


@pytest.fixture
def canonical_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "desktop.xlsx"
    workbook = Workbook()
    metadata = workbook.active
    metadata.title = "Debbie"
    metadata.append(["Key", "Value"])
    metadata.append(["Format Name", "Debbie Nesting Workbook"])
    metadata.append(["Schema Version", "1.0"])
    metadata.append(["Units", "mm"])
    works = workbook.create_sheet("Works")
    works.append([
        "Work Key", "Work Name", "Batch Multiplier", "Kerf (mm)",
        "Part Clearance (mm)", "Boundary Clearance (mm)", "Trim Left (mm)",
        "Trim Right (mm)", "Trim Top (mm)", "Trim Bottom (mm)",
    ])
    works.append(["W1", "Desktop Job", 1, 0.2, 1, 2, 3, 4, 5, 6])
    parts = workbook.create_sheet("Parts")
    parts.append([
        "Work Key", "Part Key", "Part Name", "Length (mm)", "Width (mm)",
        "Quantity", "Allow Rotation", "Drawing Number",
    ])
    parts.append(["W1", "P1", "Panel", 40, 20, 2, "Yes", "D-1"])
    stocks = workbook.create_sheet("Stocks")
    stocks.append([
        "Work Key", "Stock Key", "Stock Name", "Length (mm)", "Width (mm)", "Quantity",
    ])
    stocks.append(["W1", "S1", "Sheet", 100, 60, 1])
    workbook.save(path)
    workbook.close()
    return path
