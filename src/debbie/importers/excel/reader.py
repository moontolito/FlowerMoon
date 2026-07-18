"""Bounded, offline workbook I/O with formula/cached-value separation."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any
from xml.etree.ElementTree import ParseError, fromstring
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .constants import ImportLimits
from .diagnostics import ImportDiagnosticCode


@dataclass(frozen=True, slots=True)
class CellSnapshot:
    value: Any
    row: int
    column: int
    column_letter: str
    is_formula: bool = False
    formula_text: str | None = None
    is_error: bool = False

    @property
    def value_repr(self) -> str:
        return repr(self.formula_text if self.is_formula else self.value)


@dataclass(frozen=True, slots=True)
class MergedRangeSnapshot:
    coordinate: str
    min_row: int
    max_row: int
    min_column: int
    max_column: int


@dataclass(frozen=True, slots=True)
class SheetSnapshot:
    title: str
    rows: tuple[tuple[CellSnapshot, ...], ...]
    max_row: int
    max_column: int
    merged_ranges: tuple[MergedRangeSnapshot, ...]

    def cell(self, row: int, column: int) -> CellSnapshot:
        if row <= 0 or column <= 0:
            raise ValueError("worksheet coordinates are one-based")
        if row <= len(self.rows) and column <= len(self.rows[row - 1]):
            return self.rows[row - 1][column - 1]
        return CellSnapshot(None, row, column, get_column_letter(column))

    def row_has_value(self, row: int) -> bool:
        if row <= 0 or row > len(self.rows):
            return False
        return any(cell.value is not None or cell.is_formula for cell in self.rows[row - 1])


@dataclass(frozen=True, slots=True)
class WorkbookSnapshot:
    sheets: tuple[SheetSnapshot, ...]


class WorkbookReadProblem(Exception):
    def __init__(self, code: ImportDiagnosticCode, message: str) -> None:
        super().__init__(message)
        self.code = code


def _contains_forbidden_xml_declaration(stream) -> bool:
    """Find active DTD/entity declarations while ignoring comments and CDATA."""

    state = "normal"
    buffer = b""
    while True:
        chunk = stream.read(64 * 1024)
        final = not chunk
        # OOXML normally uses UTF-8, but removing NUL code-unit padding also
        # detects ASCII declaration tokens in UTF-16/UTF-32 XML safely.
        buffer += chunk.lower().replace(b"\x00", b"")
        while buffer:
            if state == "comment":
                end = buffer.find(b"-->")
                if end < 0:
                    buffer = b"" if final else buffer[-2:]
                    break
                buffer = buffer[end + 3 :]
                state = "normal"
                continue
            if state == "cdata":
                end = buffer.find(b"]]>")
                if end < 0:
                    buffer = b"" if final else buffer[-2:]
                    break
                buffer = buffer[end + 3 :]
                state = "normal"
                continue

            markers = {
                "comment": buffer.find(b"<!--"),
                "cdata": buffer.find(b"<![cdata["),
                "doctype": buffer.find(b"<!doctype"),
                "entity": buffer.find(b"<!entity"),
            }
            present = {name: index for name, index in markers.items() if index >= 0}
            if not present:
                buffer = b"" if final else buffer[-16:]
                break
            marker = min(present, key=present.get)
            position = present[marker]
            if marker in {"doctype", "entity"}:
                return True
            state = marker
            marker_length = 4 if marker == "comment" else 9
            buffer = buffer[position + marker_length :]
        if final:
            return False


def _check_archive(path: Path, limits: ImportLimits) -> None:
    try:
        file_size = path.stat().st_size
    except FileNotFoundError as error:
        raise WorkbookReadProblem(
            ImportDiagnosticCode.FILE_NOT_FOUND, f"Workbook does not exist: {path}"
        ) from error
    if file_size > limits.max_file_bytes:
        raise WorkbookReadProblem(
            ImportDiagnosticCode.SIZE_LIMIT_EXCEEDED,
            f"Workbook file exceeds the {limits.max_file_bytes}-byte limit",
        )
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > limits.max_archive_members:
                raise WorkbookReadProblem(
                    ImportDiagnosticCode.SIZE_LIMIT_EXCEEDED,
                    "Workbook archive contains too many members",
                )
            total = sum(item.file_size for item in members)
            if total > limits.max_archive_uncompressed_bytes:
                raise WorkbookReadProblem(
                    ImportDiagnosticCode.SIZE_LIMIT_EXCEEDED,
                    "Workbook archive expands beyond the configured limit",
                )
            for item in members:
                if item.flag_bits & 0x1:
                    raise WorkbookReadProblem(
                        ImportDiagnosticCode.UNREADABLE_WORKBOOK,
                        "Encrypted workbook archives are not supported",
                    )
                if item.file_size > limits.max_archive_member_bytes:
                    raise WorkbookReadProblem(
                        ImportDiagnosticCode.SIZE_LIMIT_EXCEEDED,
                        f"Workbook member {item.filename!r} exceeds the configured limit",
                    )
                compressed = max(item.compress_size, 1)
                if item.file_size / compressed > limits.max_archive_compression_ratio:
                    raise WorkbookReadProblem(
                        ImportDiagnosticCode.SIZE_LIMIT_EXCEEDED,
                        f"Workbook member {item.filename!r} has an unsafe compression ratio",
                    )
            for item in members:
                if not item.filename.casefold().endswith((".xml", ".rels")):
                    continue
                with archive.open(item) as stream:
                    if _contains_forbidden_xml_declaration(stream):
                        raise WorkbookReadProblem(
                            ImportDiagnosticCode.UNSAFE_WORKBOOK_CONTENT,
                            f"Workbook XML member {item.filename!r} contains "
                            "a forbidden DTD or entity declaration",
                        )
            if any(
                PurePosixPath(item.filename.replace("\\", "/")).name.casefold()
                == "vbaproject.bin"
                for item in members
            ):
                raise WorkbookReadProblem(
                    ImportDiagnosticCode.UNSUPPORTED_FILE_TYPE,
                    "Macro-enabled workbook content is not supported",
                )
            manifest_name = next(
                (
                    item.filename
                    for item in members
                    if item.filename.casefold() == "[content_types].xml"
                ),
                None,
            )
            if manifest_name is not None:
                try:
                    manifest = fromstring(archive.read(manifest_name))
                except ParseError as error:
                    raise WorkbookReadProblem(
                        ImportDiagnosticCode.UNREADABLE_WORKBOOK,
                        "The workbook content-type manifest is malformed",
                    ) from error
                macro_content_types = {
                    "application/vnd.ms-excel.sheet.macroenabled.main+xml",
                    "application/vnd.ms-excel.template.macroenabled.main+xml",
                    "application/vnd.ms-excel.addin.macroenabled.main+xml",
                    "application/vnd.ms-office.vbaproject",
                }
                if any(
                    element.attrib.get("ContentType", "").casefold()
                    in macro_content_types
                    for element in manifest.iter()
                ):
                    raise WorkbookReadProblem(
                        ImportDiagnosticCode.UNSUPPORTED_FILE_TYPE,
                        "Macro-enabled OOXML content is not supported even with an .xlsx name",
                    )
    except BadZipFile as error:
        raise WorkbookReadProblem(
            ImportDiagnosticCode.UNREADABLE_WORKBOOK,
            "The file is not a valid .xlsx ZIP package",
        ) from error


def read_workbook(path: Path, limits: ImportLimits) -> WorkbookSnapshot:
    """Read values and formula markers without retaining workbook objects.

    The data-only workbook supplies cached values. A second non-data-only read
    detects formula cells; formulas are never evaluated. External-link caches
    and VBA content are not retained.
    """

    _check_archive(path, limits)
    data_book = None
    formula_book = None
    try:
        data_book = load_workbook(
            path,
            read_only=False,
            data_only=True,
            keep_vba=False,
            keep_links=False,
        )
        formula_book = load_workbook(
            path,
            read_only=False,
            data_only=False,
            keep_vba=False,
            keep_links=False,
        )
        if len(data_book.worksheets) > limits.max_worksheets:
            raise WorkbookReadProblem(
                ImportDiagnosticCode.SIZE_LIMIT_EXCEEDED,
                f"Workbook contains more than {limits.max_worksheets} worksheets",
            )

        sheets: list[SheetSnapshot] = []
        for data_sheet in data_book.worksheets:
            formula_sheet = formula_book[data_sheet.title]
            max_row = max(data_sheet.max_row, formula_sheet.max_row)
            max_column = max(data_sheet.max_column, formula_sheet.max_column)
            if max_row > limits.max_rows_per_sheet:
                raise WorkbookReadProblem(
                    ImportDiagnosticCode.SIZE_LIMIT_EXCEEDED,
                    f"Worksheet {data_sheet.title!r} exceeds the row limit",
                )
            if max_column > limits.max_columns_per_sheet:
                raise WorkbookReadProblem(
                    ImportDiagnosticCode.SIZE_LIMIT_EXCEEDED,
                    f"Worksheet {data_sheet.title!r} exceeds the column limit",
                )

            rows: list[tuple[CellSnapshot, ...]] = []
            for row in range(1, max_row + 1):
                cells: list[CellSnapshot] = []
                for column in range(1, max_column + 1):
                    value_cell = data_sheet.cell(row, column)
                    formula_cell = formula_sheet.cell(row, column)
                    is_formula = formula_cell.data_type == "f"
                    formula_text = str(formula_cell.value) if is_formula else None
                    cells.append(
                        CellSnapshot(
                            value=value_cell.value,
                            row=row,
                            column=column,
                            column_letter=get_column_letter(column),
                            is_formula=is_formula,
                            formula_text=formula_text,
                            is_error=value_cell.data_type == "e" or formula_cell.data_type == "e",
                        )
                    )
                rows.append(tuple(cells))
            merged = tuple(
                MergedRangeSnapshot(
                    coordinate=str(cell_range),
                    min_row=cell_range.min_row,
                    max_row=cell_range.max_row,
                    min_column=cell_range.min_col,
                    max_column=cell_range.max_col,
                )
                for cell_range in formula_sheet.merged_cells.ranges
            )
            sheets.append(
                SheetSnapshot(
                    title=data_sheet.title,
                    rows=tuple(rows),
                    max_row=max_row,
                    max_column=max_column,
                    merged_ranges=merged,
                )
            )
        return WorkbookSnapshot(tuple(sheets))
    except WorkbookReadProblem:
        raise
    except (OSError, ValueError, KeyError, BadZipFile) as error:
        raise WorkbookReadProblem(
            ImportDiagnosticCode.UNREADABLE_WORKBOOK,
            f"Workbook could not be read: {type(error).__name__}",
        ) from error
    finally:
        if data_book is not None:
            data_book.close()
        if formula_book is not None:
            formula_book.close()
